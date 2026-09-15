import html
import io
import logging
import re
from collections import defaultdict
from datetime import datetime

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt

from aiogram.types import InputRichMessage

logger = logging.getLogger(__name__)


def extract_network(text: str) -> str:
    """
    Extracts radio network / unit name from message text.
    Examples:
      'Радиосеть 1 шб 22 ошбр.' -> '1 шб 22 ошбр'
      'Радиосеть БТГр 222 омбр.' -> 'БТГр 222 омбр'
      'Радиосеть 212 омбр.' -> '212 омбр'
      'Радиосеть н/у подразделения.' -> 'н/у подразделения'
    """
    if not text:
        return "н/у подразделения"

    m = re.search(
        r"(?:радиосеть|радиосети|радиосетей)\s*[:№\-]?\s*([^.\n\r]+)",
        text,
        re.IGNORECASE,
    )
    if m:
        val = m.group(1).strip()
        val = re.sub(r"[\.,;]+$", "", val).strip()
        if val:
            return val
    return "н/у подразделения"


def parse_period(text: str) -> tuple[datetime, datetime] | None:
    """
    Parses datetime period string in format: 'dd.mm.yyyy HH.MM - dd.mm.yyyy HH.MM'.
    Accepts dots, colons, underscores and hyphens/dashes.
    """
    cleaned = text.strip().strip('"\'«»')
    pattern = (
        r"(\d{2}\.\d{2}\.\d{4})\s*[_ ]\s*(\d{2}[.:]\d{2})\s*[-—–]\s*"
        r"(\d{2}\.\d{2}\.\d{4})\s*[_ ]\s*(\d{2}[.:]\d{2})"
    )
    m = re.search(pattern, cleaned)
    if not m:
        return None

    d1, t1, d2, t2 = m.groups()
    t1 = t1.replace(":", ".")
    t2 = t2.replace(":", ".")

    try:
        dt1 = datetime.strptime(f"{d1} {t1}", "%d.%m.%Y %H.%M")
        dt2 = datetime.strptime(f"{d2} {t2}", "%d.%m.%Y %H.%M")
        if dt1 > dt2:
            dt1, dt2 = dt2, dt1
        return dt1, dt2
    except ValueError:
        return None


def create_stats_report(
    period_str: str,
    messages: list[dict],
) -> InputRichMessage:
    """
    Builds the statistics report as an InputRichMessage with a bordered native table.
    Lists radio networks, message counts per network, and a 'Всего' total row.
    """
    unit_groups: dict[str, list[dict]] = defaultdict(list)
    total_count = len(messages)

    for msg in messages:
        unit = extract_network(msg["message"])
        unit_groups[unit].append(msg)

    sorted_units = sorted(
        unit_groups.items(), key=lambda item: len(item[1]), reverse=True
    )

    rows_html = []
    for unit, msgs in sorted_units:
        count = len(msgs)
        escaped_unit = html.escape(unit)
        rows_html.append(
            f"<tr><td>{escaped_unit}</td><td><code>{count}</code></td></tr>"
        )

    # Add 'Всего' summary row
    rows_html.append(
        f"<tr><td><b>Всего</b></td><td><b><code>{total_count}</code></b></td></tr>"
    )

    table_rows = "".join(rows_html)
    escaped_period = html.escape(period_str)

    html_content = (
        f"<p><b>Итого за период {escaped_period}</b></p>"
        f'<table border="1">'
        f"<thead><tr><th>Радиосеть</th><th>Радиограмм</th></tr></thead>"
        f"<tbody>{table_rows}</tbody>"
        f"</table>"
        f'<p>Чтобы получить статистику напишите период и время в формате "дд.мм.гггг чч.мм - дд.мм.гггг чч.мм".</p>'
    )

    return InputRichMessage(html=html_content)


def generate_stats_docx(
    period_str: str,
    messages: list[dict],
) -> io.BytesIO:
    """
    Generates a Microsoft Word (.docx) document with a styled table
    containing the radio network statistics for the specified period.
    """
    unit_groups: dict[str, list[dict]] = defaultdict(list)
    total_count = len(messages)

    for msg in messages:
        unit = extract_network(msg["message"])
        unit_groups[unit].append(msg)

    sorted_units = sorted(
        unit_groups.items(), key=lambda item: len(item[1]), reverse=True
    )

    doc = docx.Document()

    # Document Header / Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(f"Итого за период {period_str}")
    title_run.bold = True
    title_run.font.size = Pt(14)

    # Create Table with Borders
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # Header Row
    hdr_cells = table.rows[0].cells
    hdr_cells[0].width = Inches(3.5)
    hdr_cells[1].width = Inches(2.2)

    p0 = hdr_cells[0].paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run("Радиосеть")
    r0.bold = True

    p1 = hdr_cells[1].paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run("Количество радиограмм")
    r1.bold = True

    # Background shading for header
    for cell in hdr_cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_pr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="EAEAEA"/>'))

    # Data Rows
    for unit, msgs in sorted_units:
        count = len(msgs)
        row_cells = table.add_row().cells
        row_cells[0].width = Inches(3.5)
        row_cells[1].width = Inches(2.2)

        row_cells[0].paragraphs[0].text = unit
        row_p1 = row_cells[1].paragraphs[0]
        row_p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        row_p1.text = str(count)

    # Total Row
    tot_cells = table.add_row().cells
    tot_cells[0].width = Inches(3.5)
    tot_cells[1].width = Inches(2.2)

    tot_p0 = tot_cells[0].paragraphs[0]
    tot_r0 = tot_p0.add_run("Всего")
    tot_r0.bold = True

    tot_p1 = tot_cells[1].paragraphs[0]
    tot_p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tot_r1 = tot_p1.add_run(str(total_count))
    tot_r1.bold = True

    for cell in tot_cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_pr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F2F2F2"/>'))

    # Instruction footnote
    note_p = doc.add_paragraph()
    note_p.paragraph_format.space_before = Pt(12)
    note_run = note_p.add_run(
        'Чтобы получить статистику напишите период и время в формате "дд.мм.гггг чч.мм - дд.мм.гггг чч.мм".'
    )
    note_run.font.italic = True
    note_run.font.size = Pt(9)

    doc_stream = io.BytesIO()
    doc.save(doc_stream)
    doc_stream.seek(0)
    return doc_stream
